"""Builds the Excel template for the viewer's Excel import.

    python3 scripts/build-xlsx-template.py                      -> site/templates/boring-log-template.xlsx
    python3 scripts/build-xlsx-template.py examples/NAS-1.json  -> site/templates/NAS-1.xlsx (filled in)
    python3 scripts/build-xlsx-template.py LOG.json --bare OUT  -> no instruction rows (for tests)

The sheets, field names and explanations come from SHEETS in src/xlsx.js, so
the template always matches what the importer reads. Needs Node and openpyxl
(pip install openpyxl).
"""
import json
import os
import subprocess
import sys

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..')

SHEETS = json.loads(subprocess.check_output(
    ['node', '--input-type=module', '-e',
     "import { SHEETS, SAMPLER_TYPES } from './src/xlsx.js'; console.log(JSON.stringify({ SHEETS, SAMPLER_TYPES }))"],
    cwd=ROOT))
SAMPLER_TYPES = SHEETS['SAMPLER_TYPES']
SHEETS = SHEETS['SHEETS']

TITLE = Font(bold=True, size=14, color='1F3864')
NOTE = Font(italic=True, size=10, color='595959')
HELP = Font(size=9, color='404040')
HEAD = Font(bold=True, color='FFFFFF')
HEAD_FILL = PatternFill('solid', fgColor='2F5597')
REQ_FILL = PatternFill('solid', fgColor='C00000')
HELP_FILL = PatternFill('solid', fgColor='EEF3FA')
KEY_FILL = PatternFill('solid', fgColor='DDE6F4')
THIN = Side(style='thin', color='BFBFBF')
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP_TOP = Alignment(wrap_text=True, vertical='top')

REMOVABLE = ('The rows above the blue header row are instructions only: you can delete them. '
             'The viewer finds the header row by its field names, so keep those names as they are.')


def width_for(field):
    return {'description': 60, 'remarks': 40, 'note': 36, 'text': 60, 'url': 40, 'id': 22, 'label': 24,
            'hatch': 14, 'type': 14, 'blows': 12}.get(field['name'], max(12, len(field['name']) + 3))


def add_validation(ws, field, col, first, last=1000):
    """Drop-down lists for fields with a fixed set of values."""
    choices = None
    if field['type'] == 'choice':
        choices = field['choices']
    elif field['type'] == 'boolean':
        choices = ['TRUE', 'FALSE']
    elif ws.title == 'Samples' and field['name'] == 'type':
        choices = SAMPLER_TYPES
    if not choices:
        return
    dv = DataValidation(type='list', formula1='"' + ','.join(choices) + '"', allow_blank=True,
                        showErrorMessage=field['name'] != 'type')  # custom sampler codes are allowed
    dv.error = 'Choose one of: ' + ', '.join(choices)
    ws.add_data_validation(dv)
    dv.add(f'{col}{first}:{col}{last}')


def fields_sheet(ws, spec, values, bare):
    r = 1
    if not bare:
        ws.cell(1, 1, spec['title']).font = TITLE
        ws.cell(2, 1, spec['about']).font = NOTE
        ws.cell(3, 1, 'Fill in column B. Column C explains each field; it and these top rows can be deleted. '
                      'Keep the field names in column A as they are.').font = NOTE
        for row in (2, 3):
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
            ws.cell(row, 1).alignment = WRAP_TOP
            ws.row_dimensions[row].height = 30
        r = 5
        for c, h in enumerate(['field', 'value', 'what goes here'], 1):
            cell = ws.cell(r - 1, c, h)
            cell.font, cell.fill, cell.border = HEAD, HEAD_FILL, BOX
    for field in spec['fields']:
        key = ws.cell(r, 1, field['name'])
        key.font, key.fill, key.border = Font(bold=True), KEY_FILL, BOX
        value = ws.cell(r, 2, values.get(field['name']))
        value.border = BOX
        if field['type'] == 'date':
            value.number_format = 'yyyy-mm-dd'
        add_validation(ws, field, 'B', r, r)
        if not bare:
            h = ws.cell(r, 3, field['help'])
            h.font, h.alignment, h.border = HELP, WRAP_TOP, BOX
        r += 1
    ws.column_dimensions['A'].width = 18
    ws.column_dimensions['B'].width = 40
    ws.column_dimensions['C'].width = 80
    if not bare:
        ws.freeze_panes = 'A5'


def table_sheet(ws, spec, rows, bare, extra=()):
    fields = list(spec['fields']) + [{'name': k, 'type': 'text', 'help': 'Your own column (custom value). Show it with the Columns sheet.'} for k in extra]
    n = len(fields)
    if bare:
        head = 1
    else:
        ws.cell(1, 1, spec['title']).font = TITLE
        ws.cell(2, 1, spec['about']).font = NOTE
        ws.cell(3, 1, REMOVABLE + ' Red headers are required. Empty rows are skipped.').font = NOTE
        for row in (2, 3):
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=max(n, 6))
            ws.cell(row, 1).alignment = WRAP_TOP
            ws.row_dimensions[row].height = 30
        for c, f in enumerate(fields, 1):
            cell = ws.cell(4, c, f['help'])
            cell.font, cell.fill, cell.alignment, cell.border = HELP, HELP_FILL, WRAP_TOP, BOX
        ws.row_dimensions[4].height = 96
        head = 5
    for c, f in enumerate(fields, 1):
        cell = ws.cell(head, c, f['name'])
        cell.font, cell.border = HEAD, BOX
        cell.fill = REQ_FILL if f.get('required') else HEAD_FILL
        if bare is False:
            cell.comment = Comment(f['help'], 'Boring Log Viewer', width=300, height=120)
        col = get_column_letter(c)
        ws.column_dimensions[col].width = width_for(f)
        add_validation(ws, f, col, head + 1)
        if f['type'] == 'date':
            for r in range(head + 1, head + 200):
                ws.cell(r, c).number_format = 'yyyy-mm-dd'
    for i, row in enumerate(rows):
        for c, f in enumerate(fields, 1):
            v = row.get(f['name'], (row.get('custom') or {}).get(f['name']))
            if isinstance(v, list):
                v = ', '.join(str(x) for x in v)
            if v is not None:
                cell = ws.cell(head + 1 + i, c, v)
                cell.alignment = Alignment(wrap_text=f['name'] in ('description', 'remarks', 'text'), vertical='top')
    ws.freeze_panes = ws.cell(head + 1, 1)


def readme(ws):
    lines = [
        ('Boring log workbook', TITLE),
        ('Fill in this workbook and open it in the Boring Log Viewer (Open files…) to draw the boring log. '
         'One workbook holds one boring. The viewer turns it into the boring log JSON, which you can then edit and download.', None),
        ('', None),
        ('How it is read', Font(bold=True)),
        ('• Each sheet is one part of the log. The sheets are found by their names (Layers, Samples, …), in any order. '
         'Delete the sheets you don\'t need; this README sheet is ignored.', None),
        ('• On the table sheets (Layers, Samples, Specimens, Groundwater, Depth notes, References, Columns) put one item per row under the blue header row. '
         'The header row is found by its field names, so the title and the light blue explanation rows above it can be deleted. Columns can be in any order; '
         'columns you leave empty can be deleted. Red headers are required.', None),
        ('• On the field/value sheets (Units, Metadata, Layout) fill in column B next to each field name. Column C only explains the field.', None),
        ('• Units: every depth, diameter and unit weight is in the units chosen on the Units sheet (m, kN/m3 and mm if it is left empty). '
         'Don\'t type units into the cells: 1.5, not 1.5 m.', None),
        ('• Specimens (laboratory results) belong to a sample: give the sample\'s name in their sample_name column.', None),
        ('• Your own values: add a column with any name (letters, digits and _) to the Samples or Specimens sheet, '
         'then list it on the Columns sheet with source "sample" or "specimen" to show it on the log.', None),
        ('• Dates can be Excel dates or text such as 2024-05-01. TRUE/FALSE cells also accept yes/no.', None),
        ('', None),
        ('Sheets', Font(bold=True)),
    ]
    for s in SHEETS:
        lines.append((f'{s["title"]}: {s["about"]}', None))
    lines += [('', None), ('More help: the "Excel workbooks" page of the Boring Log Viewer documentation.', NOTE)]
    for r, (text, font) in enumerate(lines, 1):
        cell = ws.cell(r, 1, text)
        cell.alignment = Alignment(wrap_text=True, vertical='top')
        if font:
            cell.font = font
    ws.column_dimensions['A'].width = 120


def build(doc, out, bare):
    wb = Workbook()
    first = wb.active
    if bare:
        wb.remove(first)
    else:
        first.title = 'README'
        readme(first)
    samples = doc.get('samples', [])
    specimens = [{**sp, 'sample_name': s.get('name')} for s in samples for sp in s.get('specimens', [])]
    layout = doc.get('layout', {})
    columns = [c if isinstance(c, dict) else {'id': c} for c in layout.get('columns', [])]

    def extra(items):
        keys = []
        for it in items:
            for k in (it.get('custom') or {}):
                if k not in keys:
                    keys.append(k)
        return keys

    data = {
        'units': doc.get('units', {}), 'metadata': doc.get('metadata', {}), 'layout': layout,
        'layers': doc.get('layers', []), 'samples': samples, 'specimens': specimens,
        'groundwater': doc.get('groundwater', []), 'depth_notes': doc.get('depth_notes', []),
        'references': doc.get('references', []), 'columns': columns,
    }
    for spec in SHEETS:
        ws = wb.create_sheet(spec['title'])
        if spec['form'] == 'fields':
            fields_sheet(ws, spec, data[spec['key']], bare)
        else:
            table_sheet(ws, spec, data[spec['key']], bare, extra(data[spec['key']]) if spec.get('custom') else ())
    os.makedirs(os.path.dirname(os.path.abspath(out)), exist_ok=True)
    wb.save(out)
    print('wrote', out)


if __name__ == '__main__':
    args = [a for a in sys.argv[1:] if a != '--bare']
    bare = '--bare' in sys.argv
    src = args[0] if args else None
    doc = json.load(open(src, encoding='utf-8')) if src else {}
    if len(args) > 1:
        out = args[1]
    elif src:
        out = os.path.join(ROOT, 'site', 'templates', os.path.splitext(os.path.basename(src))[0] + '.xlsx')
    else:
        out = os.path.join(ROOT, 'site', 'templates', 'boring-log-template.xlsx')
    build(doc, out, bare)
